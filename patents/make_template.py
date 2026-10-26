from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
import sys
F="Arial"
thin=Side(style="thin",color="BFBFBF"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
H=PatternFill("solid",fgColor="1F4E78"); Y=PatternFill("solid",fgColor="FFF2CC"); G=PatternFill("solid",fgColor="F2F2F2")
hf=Font(name=F,bold=True,color="FFFFFF"); nf=Font(name=F,size=10)
wb=Workbook()

# 1. 특허목록
ws=wb.active; ws.title="특허목록"
cols=[("보고일",11),("업체",13),("출원인(원문)",22),("국가",7),("공개/등록번호",18),("구분",8),("제목",34),("공개일",11),("출원일",11),("설계 분야",14),("용도",12),("핵심 내용(설계 관점)",50),("링크",30),("중요도",8),("검토 메모",30)]
for i,(c,w) in enumerate(cols,1):
    x=ws.cell(1,i,c); x.font=hf; x.fill=H; x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); x.border=B
    ws.column_dimensions[x.column_letter].width=w
for i in (14,15): ws.cell(1,i).fill=PatternFill("solid",fgColor="BF8F00")
ws.row_dimensions[1].height=30
ex=["2026-11-02","LG전자","LG Electronics Inc.","KR","(예시) KR 10-2026-0000000","공개","(예시) 원심압축기의 베인 디퓨저 구조","2026-10-28","2025-04-15","디퓨저·볼류트","냉매(칠러)","(예시 행 – 실제 특허 아님) 디퓨저 베인 각도를 유량에 따라 가변해 저유량 서지 마진 확보. 기존 대비 작동 범위 확대 주장.","https://patents.google.com/...","상",""]
EMPTY=len(sys.argv)>2
for i,v in enumerate([] if EMPTY else ex,1):
    x=ws.cell(2,i,v); x.font=Font(name=F,size=10,italic=True,color="7F7F7F"); x.border=B; x.alignment=Alignment(vertical="top",wrap_text=True)
for r in range(2 if EMPTY else 3,1003):
    for i in range(1,16):
        x=ws.cell(r,i); x.font=nf; x.border=B; x.alignment=Alignment(vertical="top",wrap_text=True)
        if i in (14,15): x.fill=Y
ws.freeze_panes="C2"; ws.auto_filter.ref="A1:O1002"
ws.cell(1,10).comment=Comment("임펠러 / 디퓨저·볼류트 / 에어포일 베어링 / 기타 베어링(자기·볼 등) / 회전체 구조(로터·축·임펠러 체결·로터다이내믹스) / 모터 / 냉각유로(모터·베어링 냉각 유로) / 제어·서지 / 시스템 / 기타","Claude")
ws.cell(1,14).comment=Comment("검토자 입력: 상 / 중 / 하","Claude")
ws.cell(1,6).comment=Comment("공개 = 출원 공개, 등록 = 특허 등록","Claude")
comp=["LG전자","댄포스","코플랜드","덴소","다이킨","티앤이코리아","터보맥스","기타"]
field=["임펠러","디퓨저·볼류트","에어포일 베어링","기타 베어링","회전체 구조","모터","냉각유로","제어·서지","시스템","기타"]
lists=wb.create_sheet("목록")
for i,v in enumerate(comp,1): lists.cell(i,1,v)
for i,v in enumerate(field,1): lists.cell(i,2,v)
for i,v in enumerate(["상","중","하"],1): lists.cell(i,3,v)
for i,v in enumerate(["공개","등록"],1): lists.cell(i,4,v)
for i,v in enumerate(["냉매(칠러)","냉매(히트펌프)","공기(블로어)","차량·연료전지","기타"],1): lists.cell(i,5,v)
for i,v in enumerate(["KR","US","EP","WO","JP","CN","DE","기타"],1): lists.cell(i,6,v)
lists.sheet_state="hidden"
for col,rng in [("B","$A$1:$A$8"),("J","$B$1:$B$10"),("N","$C$1:$C$3"),("F","$D$1:$D$2"),("K","$E$1:$E$5"),("D","$F$1:$F$8")]:
    dv=DataValidation(type="list",formula1=f"=목록!{rng}",allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}2:{col}1002")

# 2. 업체별 요약
s=wb.create_sheet("업체별 요약",0)
s["A1"]="업체 × 설계 분야 누적 건수 (특허목록 시트에서 자동 집계)"; s["A1"].font=Font(name=F,bold=True,size=12)
s.cell(3,1,"업체").font=hf; s.cell(3,1).fill=H; s.cell(3,1).border=B
for j,f in enumerate(field,2):
    x=s.cell(3,j,f); x.font=hf; x.fill=H; x.border=B; x.alignment=Alignment(horizontal="center",wrap_text=True); s.column_dimensions[x.column_letter].width=12
for j,t in [(len(field)+2,"합계"),(len(field)+3,"중요도 '상'")]:
    x=s.cell(3,j,t); x.font=hf; x.fill=H; x.border=B; x.alignment=Alignment(horizontal="center"); s.column_dimensions[x.column_letter].width=12
s.column_dimensions["A"].width=14; s.row_dimensions[3].height=30
for i,c in enumerate(comp,4):
    s.cell(i,1,c).font=Font(name=F,bold=True); s.cell(i,1).border=B
    NF=len(field); TC=NF+2; HC=NF+3
    for j in range(2,NF+2):
        col=s.cell(3,j).column_letter
        x=s.cell(i,j,f'=COUNTIFS(특허목록!$B:$B,$A{i},특허목록!$J:$J,{col}$3)'); x.font=nf; x.border=B; x.alignment=Alignment(horizontal="center"); x.number_format='0;-0;"-"'
    s.cell(i,TC,f"=SUM(B{i}:{s.cell(i,NF+1).column_letter}{i})"); s.cell(i,HC,f'=COUNTIFS(특허목록!$B:$B,$A{i},특허목록!$N:$N,"상")')
    for j in (TC,HC): x=s.cell(i,j); x.font=Font(name=F,bold=True); x.border=B; x.alignment=Alignment(horizontal="center"); x.number_format='0;-0;"-"'
r=4+len(comp)
s.cell(r,1,"합계").font=Font(name=F,bold=True)
for j in range(1,len(field)+4):
    x=s.cell(r,j); x.fill=G; x.border=B
    if j>1:
        L=x.column_letter; x.value=f"=SUM({L}4:{L}{r-1})"; x.font=Font(name=F,bold=True); x.alignment=Alignment(horizontal="center"); x.number_format='0;-0;"-"'

# 3. 사용법
u=wb.create_sheet("사용법")
u.column_dimensions["A"].width=100
lines=["터보(원심) 압축기 특허 주간 정리 – 엑셀 양식 (검토용 초안)","",
"■ 시트 구성",
"  · 업체별 요약 : 업체 × 설계 분야 건수가 자동으로 계산됩니다 (직접 입력 X).",
"  · 특허목록   : 매주 새 특허가 아래로 계속 쌓입니다. 1주 1파일이 아니라 누적 방식.",
"",
"■ 누가 무엇을 채우나",
"  · A~M열 (파란 머리글) : Claude가 매주 자동 입력.",
"  · N~O열 (노란 칸, 금색 머리글) : 검토자(본인)가 중요도(상/중/하)와 메모를 입력.",
"",
"■ 2행은 형식을 보여주는 예시입니다 (실제 특허 아님). 실제 운영 시 삭제됩니다.",
"",
"■ 드롭다운(선택 목록) 칸 : 업체, 국가, 구분, 설계 분야, 용도, 중요도",
"■ 머리글에 필터가 걸려 있어 업체·설계 분야·중요도로 걸러 볼 수 있습니다.",
"",
"■ 검토해 주실 점",
"  1) 열(항목)이 빠졌거나 필요 없는 것이 있는지",
"  2) 설계 분야 분류(임펠러/디퓨저·볼류트/에어포일 베어링/기타 베어링/회전체 구조/모터/냉각유로/제어·서지/시스템/기타)가 적절한지",
"  ※ 중점 분야(에어포일 베어링·회전체 구조·냉각유로)는 7개 업체 외 출원인 특허도 '기타' 업체로 포함됩니다.",
"  3) 누적 1파일 방식 vs 주차별 별도 파일 방식 중 선호"]
for i,t in enumerate(lines,1):
    x=u.cell(i,1,t); x.font=Font(name=F,bold=(i==1 or t.startswith("■")),size=12 if i==1 else 10)
for w in wb.worksheets:
    w.sheet_view.zoomScale=100
wb.active=0
wb.save(sys.argv[1])
