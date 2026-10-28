"""Build a compact Excel for the weekly mail attachment from the cumulative file.

usage: python3 patents/make_mail_xlsx.py <out.xlsx> [보고일]
보고일을 주면 그 날짜 행만, 없으면 누적 전체 행을 담는다.
"""
import sys
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

SRC = "patents/특허_누적.xlsx"
out = sys.argv[1]
only = sys.argv[2] if len(sys.argv) > 2 else None

src = load_workbook(SRC)
L, S = src["특허목록"], src["업체별 요약"]
rows = [r for r in L.iter_rows(min_row=2, values_only=True) if r[0]]
if only:
    rows = [r for r in rows if str(r[0]) == only]
latest = max((str(r[0]) for r in rows), default="")

wb = Workbook()
s = wb.active
s.title = "업체별 요약"
for r in S.iter_rows(min_row=1, max_row=12):
    for c in r:
        s.cell(c.row, c.column, c.value)
l = wb.create_sheet("특허목록")
for c in L[1]:
    l.cell(1, c.column, c.value)
for i, r in enumerate(rows, 2):
    for j, v in enumerate(r[:19], 1):
        l.cell(i, j, v)

H = PatternFill("solid", fgColor="1F4E78")
G = PatternFill("solid", fgColor="E2EFDA")
Y = PatternFill("solid", fgColor="FFF2CC")
hf = Font(name="Arial", bold=True, color="FFFFFF")
nf = Font(name="Arial", size=10)
widths = [11, 13, 22, 7, 18, 8, 34, 11, 11, 14, 12, 50, 30, 8, 30, 11, 24, 48, 40]
n = len(rows) + 1
for i, w in enumerate(widths, 1):
    l.column_dimensions[l.cell(1, i).column_letter].width = w
    h = l.cell(1, i)
    h.font = hf
    h.fill = PatternFill("solid", fgColor="BF8F00") if i in (14, 15) else H
    h.alignment = Alignment(wrap_text=True, vertical="center")
    for r in range(2, n + 1):
        c = l.cell(r, i)
        c.font = nf
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i in (14, 15):
            c.fill = Y
        elif i < 14 and str(l.cell(r, 1).value) == latest:
            c.fill = G
l.freeze_panes = "C2"
l.auto_filter.ref = f"A1:S{n}"
s["A1"].font = Font(name="Arial", bold=True, size=12)
s.column_dimensions["A"].width = 14
for j in range(1, 14):
    c = s.cell(3, j)
    c.font = hf
    c.fill = H
    c.alignment = Alignment(wrap_text=True, horizontal="center")
    if j > 1:
        s.column_dimensions[c.column_letter].width = 12
for r in range(4, 13):
    for j in range(1, 14):
        s.cell(r, j).font = Font(name="Arial", bold=(j == 1 or j > 11 or r == 12))
wb.save(out)
print(out, len(rows), "rows")
